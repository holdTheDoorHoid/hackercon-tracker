"""Export: the original spreadsheet layout (plus extras) and an iCalendar feed."""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from . import db as dbm
from .scoring import compute_score, grade

XLSX_COLUMNS = [
    "Con", "Training Start Date", "Con Start Date", "Con End Date", "Dates Status", "Location City State", "Country",
    "Attendance", "Score", "Our Status", "Email / Outreach", "IOT Vill", "CFP Deadline", "CFP Status", "Trainings Deadline",
    "Training Status", "CFV Deadline", "CFV Status", "Vendor/Sponsor Deadline", "Vendor Status", "Vending", "Table Cost",
    "Hotel", "Travel", "PTO", "Website", "Contact Email", "Discord", "CFP Link", "Sponsor Docs", "Twitter/X", "Notes",
]


def _dl(deadlines, kind):
    for d in deadlines:
        if d["kind"] == kind:
            return d
    return None


def export_xlsx(years: list[int] | None = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with dbm.db() as con:
        settings = dbm.get_settings(con)
        confs = dbm.rows(con, "SELECT * FROM conferences WHERE archived=0 ORDER BY name")
        all_years = years or [r["year"] for r in dbm.rows(con, "SELECT DISTINCT year FROM editions ORDER BY year")]
        for year in all_years:
            ws = wb.create_sheet(str(year))
            ws.append(XLSX_COLUMNS)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F3A5F")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "B2"
            eds = dbm.rows(con, "SELECT e.*, c.slug FROM editions e JOIN conferences c ON c.id=e.conference_id WHERE e.year=? AND c.archived=0 ORDER BY COALESCE(e.start_date, '9999'), c.name", (year,))
            by_slug = {c["slug"]: c for c in confs}
            for e in eds:
                c = by_slug.get(e["slug"])
                if not c:
                    continue
                dls = dbm.rows(con, "SELECT * FROM deadlines WHERE edition_id=?", (e["id"],))
                editions = dbm.rows(con, "SELECT * FROM editions WHERE conference_id=?", (c["id"],))
                sc = compute_score(c, e, dls, editions, settings)
                outreach = dbm.rows(con, "SELECT * FROM outreach WHERE conference_id=? AND (edition_id=? OR edition_id IS NULL) ORDER BY date", (c["id"], e["id"]))
                cfp, tr, cfv, vend = _dl(dls, "cfp"), _dl(dls, "training"), _dl(dls, "cfv"), (_dl(dls, "vendor") or _dl(dls, "sponsor"))
                loc = ", ".join(p for p in (c.get("city"), c.get("region")) if p) or ("Online" if c.get("online") else "")
                rowv = [
                    e.get("label") or f"{c['name']} {year}",
                    _d(e.get("training_start")), _d(e.get("start_date")), _d(e.get("end_date")), e.get("dates_status"),
                    loc, c.get("country"), c.get("attendance_est") or c.get("attendance_band"),
                    f"{sc['score']} ({grade(sc['score'])})", e.get("our_status"),
                    "; ".join(f"{o['date']}: {o.get('summary') or o.get('subject') or o['channel']}" for o in outreach),
                    "Yes" if c.get("iot_village_partner") else None,
                    _d(cfp and cfp.get("due_date")), cfp and cfp.get("status"),
                    _d(tr and tr.get("due_date")), tr and tr.get("status"),
                    _d(cfv and cfv.get("due_date")), cfv and cfv.get("status"),
                    _d(vend and vend.get("due_date")), vend and vend.get("status"),
                    c.get("vending_policy") if (c.get("vending_policy") or "unknown") != "unknown" else None,
                    c.get("table_cost"), e.get("hotel"), e.get("travel"), e.get("pto"),
                    c.get("website"), c.get("contact_email"), c.get("discord_url"), c.get("cfp_url"), c.get("sponsor_docs_url"),
                    c.get("twitter"), " | ".join(x for x in (c.get("vending_notes"), e.get("notes"), c.get("notes")) if x) or None,
                ]
                ws.append(rowv)
            for i, col in enumerate(XLSX_COLUMNS, start=1):
                width = 34 if col in ("Con", "Website", "Email / Outreach", "Notes", "Sponsor Docs", "CFP Link") else 14
                ws.column_dimensions[get_column_letter(i)].width = width
            for row_cells in ws.iter_rows(min_row=2):
                for cell in row_cells:
                    if isinstance(cell.value, date):
                        cell.number_format = "yyyy-mm-dd"
        # Contacts sheet
        ws = wb.create_sheet("Contacts")
        ws.append(["Conference", "Name", "Role", "Email", "Phone", "Notes"])
        for r in dbm.rows(con, "SELECT ct.*, c.name AS cname FROM contacts ct JOIN conferences c ON c.id=ct.conference_id ORDER BY c.name"):
            ws.append([r["cname"], r.get("name"), r.get("role"), r.get("email"), r.get("phone"), r.get("notes")])
        for c in ws[1]:
            c.font = Font(bold=True)
        # Sources sheet
        ws = wb.create_sheet("Sources")
        ws.append(["Sources", "Notes"])
        for r in dbm.rows(con, "SELECT * FROM sources ORDER BY id"):
            ws.append([r.get("url"), r.get("notes")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _d(iso):
    if not iso:
        return None
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except ValueError:
        return iso


def export_ics(include_deadlines: bool = True, only_mine: bool = False) -> str:
    """All-day events for every edition with dates, plus deadlines."""
    from icalendar import Calendar, Event

    cal = Calendar()
    cal.add("prodid", "-//Hackercon Tracker//EN")
    cal.add("version", "2.0")
    cal.add("x-wr-calname", "Conferences")
    with dbm.db() as con:
        settings = dbm.get_settings(con)
        q = ("SELECT e.*, c.name AS cname, c.slug, c.city AS ccity, c.region AS cregion, c.country AS ccountry, c.website, c.online, c.lat, c.lon, "
             "c.attendance_est, c.attendance_band, c.vending_policy, c.table_cost, c.workshop_track, c.village_hosting, "
             "c.iot_village_partner, c.priority, c.past_rating FROM editions e JOIN conferences c ON c.id=e.conference_id "
             "WHERE e.start_date IS NOT NULL AND c.archived=0")
        if only_mine:
            q += " AND e.our_status IN ('applied','accepted','attending','attended')"
        for e in dbm.rows(con, q):
            e["city"] = e.get("city") or e.get("ccity")
            e["region"], e["country"] = e["cregion"], e["ccountry"]
            ev = Event()
            start = date.fromisoformat(e["start_date"])
            end = date.fromisoformat(e["end_date"]) if e.get("end_date") else start
            ev.add("uid", f"edition-{e['id']}@hackercon-tracker")
            title = e.get("label") or f"{e['cname']} {e['year']}"
            if e.get("dates_status") in ("tentative", "estimated"):
                title += " (tentative)"
            ev.add("summary", title)
            ev.add("dtstart", start)
            ev.add("dtend", end + timedelta(days=1))
            loc = ", ".join(p for p in (e.get("city"), e.get("region"), e.get("country")) if p)
            if loc:
                ev.add("location", loc)
            dls = dbm.rows(con, "SELECT * FROM deadlines WHERE edition_id=?", (e["id"],))
            eds = dbm.rows(con, "SELECT * FROM editions WHERE conference_id=?", (e["conference_id"],))
            sc = compute_score(e, e, dls, eds, settings)
            desc = f"Score {sc['score']} ({grade(sc['score'])}). Status: {e.get('our_status')}. Dates: {e.get('dates_status')}."
            if e.get("website"):
                desc += f"\n{e['website']}"
                ev.add("url", e["website"])
            ev.add("description", desc)
            cal.add_component(ev)
            if include_deadlines:
                for d in dls:
                    if not d.get("due_date") or d.get("status") in ("done", "na", "closed", "accepted", "rejected", "declined"):
                        continue
                    dv = Event()
                    dv.add("uid", f"deadline-{d['id']}@hackercon-tracker")
                    dv.add("summary", f"DEADLINE {d['kind'].upper()}: {e['cname']} {e['year']}")
                    dd = date.fromisoformat(d["due_date"])
                    dv.add("dtstart", dd)
                    dv.add("dtend", dd + timedelta(days=1))
                    if d.get("url"):
                        dv.add("url", d["url"])
                    dv.add("description", d.get("notes") or "")
                    cal.add_component(dv)
    return cal.to_ical().decode()
