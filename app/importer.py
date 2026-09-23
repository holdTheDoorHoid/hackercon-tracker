"""Import a spreadsheet in the "one sheet per year" layout and
the curated seed JSON into the database.
"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

from . import db as dbm
from .geo import geocode, split_location, norm_country

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}

STATUS_MAP = {
    "open": "open", "closed": "closed", "submitted": "submitted", "accepted": "accepted",
    "rejected": "rejected", "declined": "declined", "waitlist": "waitlist", "done": "done",
}

# Spreadsheet spellings -> canonical slug. Anything not here is slugified.
SLUG_OVERRIDES = {
    "bsidesnepa": "bsides-nepa", "bsides ct": "bsides-ct", "pancakescon": "pancakescon",
    "jawncon": "jawncon", "jawn con": "jawncon", "bsides nyc": "bsides-nyc", "saint con": "saintcon",
    "saintcon": "saintcon", "pumpcon": "pumpcon", "bsideschicago": "bsides-chicago",
    "bsides chicago": "bsides-chicago", "the long conn": "the-long-con", "bsides charleston": "bsides-charleston",
    "carolinacon": "carolinacon", "bsides delaware": "bsides-delaware", "bsides philly": "bsides-philly",
    "bsidesphilly": "bsides-philly", "nahamcon": "nahamcon", "districtcon": "districtcon",
    "cactuscon": "cactuscon", "wild west hackin' fest @ mile high": "wwhf-mile-high",
    "wild west hackin’ fest @ mile high": "wwhf-mile-high", "bsides galway": "bsides-galway",
    "bsidesics/ot": "bsides-ics-ot", "bsides seattle": "bsides-seattle", "[un]prompted": "unprompted",
    "women in cybersecurity": "wicys", "bsides limburg": "bsides-limburg", "bsides reykjavik": "bsides-reykjavik",
    "bsides regina": "bsides-regina", "isdfs": "isdfs", "dakota con": "dakotacon", "bsides roc": "bsides-roc",
    "bsides sf": "bsides-sf", "hack the bay": "hack-the-bay", "bsides lancashire": "bsides-lancashire",
    "bsides clt": "bsides-clt", "cyphercon": "cyphercon", "bsides mke": "bsides-milwaukee",
    "bsides sd": "bsides-san-diego", "bsides ok": "bsides-ok", "kernel con": "kernelcon",
    "bsides slc": "bsides-slc", "bsidesot uk": "bsides-ot-uk", "bsides groningen": "bsides-groningen",
    "bsides south jersey": "bsides-south-jersey", "bsides prague": "bsides-prague", "bsides kc": "bsides-kc",
    "bsides charm": "bsidescharm", "defcon singapore": "defcon-singapore", "bsides bud": "bsides-budapest",
    "counter spy": "counterspy", "bsides luxembourg": "bsides-luxembourg", "hack space con": "hackspacecon",
    "bsidessouthflorida": "bsides-south-florida", "notacon": "notacon", "thotcon": "thotcon",
    "bsidesnola": "bsides-nola", "nolacon": "nolacon", "bsides nash": "bsides-nashville",
    "cackalackycon": "cackalackycon", "northsec": "northsec", "bsides312": "bsides312",
    "hack miami": "hackmiami", "bsides birmingham": "bsides-birmingham", "bsides tampa": "bsides-tampa",
    "ozcon": "ozcon", "ieee symposium on security and privacy": "ieee-sp", "ekoparty miami": "ekoparty-miami",
    "bsidesknoxville": "bsides-knoxville", "death con": "deathcon", "layerone": "layerone",
    "bsides dayton": "bsides-dayton", "bsides london": "bsides-london", "hardwear.io": "hardwear-io-usa",
    "security fest": "security-fest", "big sky cyber summit": "big-sky-cyber-summit", "bsides hbg": "bsides-hbg",
    "bsides detroit": "bsides-detroit", "bsidesmaine": "bsides-maine", "diana initiative": "diana-initiative",
    "naclcon": "naclcon", "secretcon": "secretcon", "bsides kristiansand": "bsides-kristiansand",
    "bsides roanoke": "bsides-roanoke", "el bsides": "el-bsides", "bsidesportland": "bsides-pdx",
    "layer 8": "layer8", "sleuthcon": "sleuthcon", "bsidesalbany": "bsides-albany", "bsides buffalo": "bsides-buffalo",
    "bsidesfortwayne": "bsides-fort-wayne", "rvasec": "rvasec", "silicon valley cybersecurity conference": "svcc",
    "bsides boulder": "bsides-boulder", "bsides leeds": "bsides-leeds", "bsidessatx": "bsides-satx",
    "area41 security conference": "area41", "rmisc": "rmisc", "toor camp": "toorcamp", "bsides tlv": "bsides-tlv",
    "spiceworld": "spiceworld", "boardwalk bytes": "boardwalk-bytes", "open sauce": "open-sauce",
    "bsidesumeå": "bsides-umea", "bsideszadar": "bsides-zadar", "bsidesedmonton": "bsides-edmonton",
    "bsides edmonton": "bsides-edmonton", "bsidesmarrakesh": "bsides-marrakesh", "bsidesseasides": "bsides-seasides",
    "bsidesottawa": "bsides-ottawa", "bsides las": "bsides-las-vegas", "black hat usa": "black-hat-usa",
    "defcon": "def-con", "hope": "hope", "hack glasgow": "hack-glasgow", "sec-t": "sec-t",
    "blue team con": "blue-team-con", "bsides frankfurt": "bsides-frankfurt", "labscon": "labscon",
    "sec health": "sechealth", "balkan computer congress": "balccon", "bsidesmontreal": "bsides-montreal",
    "grrcon": "grrcon", "hammercon": "hammercon", "corncon": "corncon", "bsidesatl": "bsides-atlanta",
    "bsidesbozeman": "bsides-bozeman", "ics cyber security conference": "ics-cyber-security-conference",
    "c0c0n": "c0c0n", "infosecworld": "infosec-world", "bsidestc": "bsides-tc", "hack.lu": "hack-lu",
    "rstcon": "rstcon", "bsidesaugusta": "bsides-augusta", "bsidescos": "bsides-cos", "wiccon": "wiccon",
    "bsidesswfl": "bsides-swfl", "unlock your brain harden your system": "uybhys", "queen city con": "queen-city-con",
    "bsides berlin": "bsides-berlin", "bsides ics/ot": "bsides-ics-ot",
}


def slugify(name: str) -> str:
    n = name.strip().lower()
    n = re.sub(r"\s+", " ", n)
    n = n.replace("wild west hackin’ fest", "wild west hackin' fest")
    n = re.sub(r"\b(20\d\d)\b", "", n).strip()          # drop years
    n = re.sub(r"(20\d\d)$", "", n).strip()               # "WICCON2026"
    if n in SLUG_OVERRIDES:
        return SLUG_OVERRIDES[n]
    n = re.sub(r"\s+\d{1,2}\s*$", "", n).strip()          # trailing edition numbers ("RVAsec 15")
    if n in SLUG_OVERRIDES:
        return SLUG_OVERRIDES[n]
    if n.startswith("maker faire"):
        return "maker-faire-" + re.sub(r"[^a-z0-9]+", "-", n[len("maker faire"):]).strip("-")
    s = re.sub(r"[^a-z0-9]+", "-", n).strip("-")
    return SLUG_OVERRIDES.get(s, s)


def fuzzy_dates(text, year: int) -> tuple[str | None, str | None]:
    """'Oct 10-11', 'sept 5', 'sept21', 'Oct 31 nov 1', '10-Apr', datetime -> (start, end) ISO."""
    if text is None or text == "":
        return None, None
    if isinstance(text, (datetime, date)):
        d = text.date() if isinstance(text, datetime) else text
        return d.isoformat(), None
    s = str(text).strip().lower().replace(".", "")
    # "10-apr" style
    m = re.match(r"^(\d{1,2})-([a-z]{3,})$", s)
    if m:
        mo = MONTHS.get(m.group(2)[:3])
        return (date(year, mo, int(m.group(1))).isoformat(), None) if mo else (None, None)
    # "oct 31 nov 1"
    m = re.match(r"^([a-z]{3,})\s*(\d{1,2})\s*(?:-|–|to)?\s*([a-z]{3,})\s*(\d{1,2})$", s)
    if m:
        m1, m2 = MONTHS.get(m.group(1)[:3]), MONTHS.get(m.group(3)[:3])
        if m1 and m2:
            return date(year, m1, int(m.group(2))).isoformat(), date(year, m2, int(m.group(4))).isoformat()
    # "oct 10-11" / "nov 8-9" / "sept 5" / "sept21"
    m = re.match(r"^([a-z]{3,})\s*(\d{1,2})(?:\s*(?:-|–|to)\s*(\d{1,2}))?(?:st|nd|rd|th)?,?\s*(\d{4})?$", s)
    if m:
        mo = MONTHS.get(m.group(1)[:3])
        if mo:
            y = int(m.group(4)) if m.group(4) else year
            start = date(y, mo, int(m.group(2))).isoformat()
            end = date(y, mo, int(m.group(3))).isoformat() if m.group(3) else None
            return start, end
    # "march 20th, 2026" / "april 17th"
    m = re.match(r"^([a-z]{3,})\s+(\d{1,2})(?:st|nd|rd|th)?,?\s*(\d{4})?$", s)
    if m:
        mo = MONTHS.get(m.group(1)[:3])
        if mo:
            y = int(m.group(3)) if m.group(3) else year
            return date(y, mo, int(m.group(2))).isoformat(), None
    iso = dbm.parse_date(text)
    return iso, None


def fix_year(iso: str | None, sheet_year: int, kind: str) -> str | None:
    """The sheet has a lot of dates typed with the wrong year (Excel defaulted to
    the current year). Con dates always belong to the sheet year; deadlines in
    Jan-Aug of the previous year are almost certainly this year's."""
    if not iso:
        return None
    y, m, d = (int(x) for x in iso.split("-"))
    if kind == "event":
        if y != sheet_year and (y == sheet_year - 1 or y > sheet_year + 1):
            y = sheet_year
    else:  # deadline
        if y == sheet_year - 1 and m <= 8:
            y = sheet_year
        elif y > sheet_year + 1:
            y = sheet_year
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return date(y, m, 28).isoformat()


def parse_outreach(text: str | None, year: int) -> dict | None:
    if not text:
        return None
    s = str(text).strip()
    low = s.lower()
    m = re.search(r"(\d{1,2})/(\d{1,2})", s)
    when = date(year, int(m.group(1)), int(m.group(2))).isoformat() if m else None
    channel = "gmail"
    if "web" in low or "form" in low or "page" in low:
        channel = "web-form"
    if "meeting" in low:
        channel = "meeting"
    status = "meeting" if channel == "meeting" else "sent"
    return {"date": when or f"{year}-01-01", "channel": channel, "summary": s if m else f"{s} (date not recorded)", "status": status,
            "date_known": bool(m)}


def parse_workbook(path: str | Path) -> list[dict]:
    """Return a list of conference dicts (with nested editions/deadlines/outreach)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    confs: dict[str, dict] = {}
    for ws in wb.worksheets:
        if not ws.title.strip().isdigit():
            continue
        year = int(ws.title.strip())
        hdr = [(c.value or "").strip() if isinstance(c.value, str) else (c.value or "") for c in ws[1]]
        hdr = [str(h) for h in hdr]
        for r in ws.iter_rows(min_row=2, values_only=True):
            rec = {hdr[i]: v for i, v in enumerate(r) if i < len(hdr) and v not in (None, "")}
            name = rec.get("Con") or rec.get("Column 1")
            if not name or not isinstance(name, str) or not name.strip():
                continue
            name = name.strip()
            slug = slugify(name)
            loc_txt = str(rec.get("Location City State") or "")
            if slug == "bsides-london" and ("ON" in loc_txt.split(",")[-1].strip().upper().split() or str(rec.get("Country") or "").upper() in ("CA", "CANADA")):
                slug = "bsides-london-on"
                name = "BSides London (Ontario)"

            display = re.sub(r"\s+20\d\d\s*$", "", name).strip()
            c = confs.setdefault(slug, {"slug": slug, "name": display, "aliases": set(), "editions": {}, "source": "spreadsheet"})
            c["aliases"].add(name)
            if len(display) > len(c["name"]) and display.lower().startswith("bsides"):
                c["name"] = display
            loc = rec.get("Location City State")
            if loc:
                city, region, country = split_location(str(loc))
                c.setdefault("city", city)
                c.setdefault("region", region)
                if country:
                    c.setdefault("country", country)
                if str(loc).strip().lower() == "online":
                    c["online"] = 1
                    c["city"] = None
            if rec.get("Country"):
                c["country"] = norm_country(str(rec["Country"])) or c.get("country")
            att = rec.get("Attendancce") or rec.get("Attendance")
            if att is not None:
                if isinstance(att, (int, float)):
                    c["attendance_est"] = int(att)
                else:
                    c["attendance_band"] = str(att).strip()
            if str(rec.get("IOT Vill", "")).strip().lower() == "yes":
                c["iot_village_partner"] = 1
            web = rec.get("Column 4")
            if web and isinstance(web, str) and web.startswith("http"):
                c.setdefault("website", web.strip())
            sd = rec.get("Sponsor Docs") or rec.get("sponsor docs")
            if sd and isinstance(sd, str) and sd.startswith("http"):
                c["sponsor_docs_url"] = sd.strip()
            vend = rec.get("Vending")
            if vend and str(vend).strip().upper() != "NA":
                v = str(vend).strip()
                c["vending_notes"] = v
                if "no" in v.lower():
                    c["vending_policy"] = "no"
                elif "yes" in v.lower():
                    c["vending_policy"] = "yes"
            ed = c["editions"].setdefault(year, {"year": year, "label": name if re.search(r"20\d\d", name) else None,
                                                   "deadlines": [], "outreach": [], "source": "spreadsheet"})
            s, e = fuzzy_dates(rec.get("Con Start Date"), year)
            s2, e2 = fuzzy_dates(rec.get("Con End Date"), year)
            ed["start_date"] = fix_year(s, year, "event")
            ed["end_date"] = fix_year(e or s2, year, "event")
            ts, _ = fuzzy_dates(rec.get("Training Start Date"), year)
            ed["training_start"] = fix_year(ts, year, "event")
            if ed["start_date"] and ed["end_date"] and ed["end_date"] < ed["start_date"]:
                ed["start_date"], ed["end_date"] = ed["end_date"], ed["start_date"]
            for kind, dcol, scol in (("cfp", "CFP Deadline", "CFP Status"), ("training", "Trainings Deadline", "Training Status"),
                                     ("cfv", "CFV Deadline", "CFV Status")):
                dv, sv = rec.get(dcol), rec.get(scol)
                if dv is None and sv is None:
                    continue
                due, _ = fuzzy_dates(dv, year) if dv is not None and str(dv).strip().lower() not in ("closed", "no deadline", "open") else (None, None)
                due = fix_year(due, year, "deadline")
                status = STATUS_MAP.get(str(sv).strip().lower(), None) if sv else None
                if status is None and isinstance(dv, str) and dv.strip().lower() in STATUS_MAP:
                    status = STATUS_MAP[dv.strip().lower()]
                notes = None
                if isinstance(dv, str) and dv.strip().lower() == "no deadline":
                    notes = "Rolling / no deadline"
                ed["deadlines"].append({"kind": kind, "due_date": due, "status": status or ("open" if due else "open"), "notes": notes})
            for col in ("Hotel", "Travel", "PTO"):
                if rec.get(col) is not None:
                    ed[col.lower()] = str(rec[col]).strip()
            o = parse_outreach(rec.get("Email"), year)
            if o:
                ed["outreach"].append(o)
            # a very rough "our status" from the CFP/CFV columns
            statuses = {d["status"] for d in ed["deadlines"]}
            if "accepted" in statuses:
                ed["our_status"] = "accepted"
            elif "submitted" in statuses:
                ed["our_status"] = "applied"
            elif "rejected" in statuses or "declined" in statuses:
                ed["our_status"] = "rejected" if "rejected" in statuses else "declined"
    out = []
    for c in confs.values():
        c["aliases"] = sorted(c["aliases"])
        c["editions"] = [c["editions"][y] for y in sorted(c["editions"])]
        if c["slug"].startswith("maker-faire"):
            c["type"] = "maker"
        elif c["slug"].startswith("bsides"):
            c["type"] = "bsides"
        out.append(c)
    return out


def upsert_conference(con, c: dict, overwrite: bool = False) -> int:
    """Insert or update a conference (and its editions/deadlines/outreach/contacts) from a dict."""
    existing = dbm.row(con, "SELECT * FROM conferences WHERE slug=?", (c["slug"],))
    fields = {k: v for k, v in c.items() if k in CONF_COLS and k not in ("id",)}
    if "aliases" in fields and not isinstance(fields["aliases"], str):
        fields["aliases"] = json.dumps(sorted(set(fields["aliases"])))
    if "date_hints" in fields and not isinstance(fields["date_hints"], str):
        fields["date_hints"] = json.dumps(fields["date_hints"] or [])
    if fields.get("lat") is None and (fields.get("city") or c.get("city")) and not fields.get("online"):
        g = geocode(fields.get("city") or c.get("city"), fields.get("region") or c.get("region"), fields.get("country") or c.get("country"))
        if g:
            fields["lat"], fields["lon"] = g
    if existing:
        cid = existing["id"]
        if overwrite:
            dbm.update_row(con, "conferences", cid, fields)
        else:
            # fill blanks only
            upd = {k: v for k, v in fields.items() if v not in (None, "", "unknown", 0) and existing.get(k) in (None, "", "unknown", 0, "[]")}
            if "aliases" in fields:
                merged = sorted(set(dbm.jloads(existing.get("aliases"))) | set(json.loads(fields["aliases"])))
                upd["aliases"] = json.dumps(merged)
            if upd:
                dbm.update_row(con, "conferences", cid, upd)
    else:
        fields.setdefault("name", c["slug"])
        cid = dbm.insert_row(con, "conferences", fields)
    for e in c.get("editions", []):
        upsert_edition(con, cid, e, overwrite)
    for ct in c.get("contacts", []):
        if ct.get("email") and not dbm.row(con, "SELECT id FROM contacts WHERE conference_id=? AND email=?", (cid, ct["email"])):
            dbm.insert_row(con, "contacts", {"conference_id": cid, **{k: ct.get(k) for k in ("name", "role", "email", "phone", "notes", "source")}})
    return cid


def upsert_edition(con, cid: int, e: dict, overwrite: bool = False) -> int:
    ex = dbm.row(con, "SELECT * FROM editions WHERE conference_id=? AND year=?", (cid, e["year"]))
    fields = {k: v for k, v in e.items() if k in EDITION_COLS}
    if ex:
        eid = ex["id"]
        if overwrite:
            dbm.update_row(con, "editions", eid, fields)
        else:
            upd = {k: v for k, v in fields.items() if v not in (None, "") and ex.get(k) in (None, "", "unknown", "considering")}
            if upd:
                dbm.update_row(con, "editions", eid, upd)
    else:
        eid = dbm.insert_row(con, "editions", {"conference_id": cid, **fields})
    for d in e.get("deadlines", []):
        exd = dbm.row(con, "SELECT * FROM deadlines WHERE edition_id=? AND kind=? AND (label IS ? OR label=?)", (eid, d["kind"], d.get("label"), d.get("label")))
        df = {k: v for k, v in d.items() if k in DEADLINE_COLS}
        if exd:
            upd = {k: v for k, v in df.items() if v not in (None, "") and (overwrite or exd.get(k) in (None, "", "open"))}
            if upd:
                dbm.update_row(con, "deadlines", exd["id"], upd)
        else:
            dbm.insert_row(con, "deadlines", {"edition_id": eid, **df})
    for o in e.get("outreach", []):
        if not dbm.row(con, "SELECT id FROM outreach WHERE conference_id=? AND date=? AND summary=?", (cid, o["date"], o.get("summary"))):
            dbm.insert_row(con, "outreach", {"conference_id": cid, "edition_id": eid, "date": o["date"], "channel": o.get("channel", "gmail"),
                                             "direction": o.get("direction", "out"), "subject": o.get("subject"), "summary": o.get("summary"),
                                             "status": o.get("status", "sent"), "follow_up_due": o.get("follow_up_due")})
    return eid


CONF_COLS = {
    "slug", "name", "aliases", "type", "website", "city", "region", "country", "lat", "lon", "online", "typical_month",
    "typical_pattern", "attendance_est", "attendance_band", "description", "contact_email", "cfp_email", "sponsor_email",
    "discord_url", "twitter", "mastodon", "bluesky", "linkedin", "instagram", "youtube", "cfp_url", "cfv_url", "sponsor_url",
    "sponsor_docs_url", "vending_policy", "vending_notes", "table_cost", "workshop_track", "village_hosting",
    "iot_village_partner", "ticket_price", "relationship_notes", "past_rating", "past_sales", "past_notes", "priority",
    "archived", "source", "last_verified", "verification_notes", "notes", "interval_years", "date_hints", "hints_checked",
}
EDITION_COLS = {
    "year", "label", "start_date", "end_date", "training_start", "training_end", "dates_status", "venue", "city", "url",
    "attendance_actual", "our_status", "hotel", "travel", "pto", "budget_est", "sales", "leads", "worth_it", "aar_notes",
    "notes", "source", "last_verified",
}
DEADLINE_COLS = {"kind", "label", "due_date", "status", "url", "notes"}


def import_xlsx(path, overwrite: bool = False) -> dict:
    confs = parse_workbook(path)
    n_new = n_upd = 0
    with dbm.db() as con:
        for c in confs:
            before = dbm.row(con, "SELECT id FROM conferences WHERE slug=?", (c["slug"],))
            upsert_conference(con, c, overwrite)
            if before:
                n_upd += 1
            else:
                n_new += 1
        dbm.log_activity(con, "import", f"Imported spreadsheet {Path(path).name}: {n_new} new, {n_upd} updated")
    return {"new": n_new, "updated": n_upd, "total": len(confs)}


def import_seed(path, overwrite: bool = False) -> dict:
    data = json.loads(Path(path).read_text())
    confs = data["conferences"] if isinstance(data, dict) else data
    n_new = n_upd = 0
    with dbm.db() as con:
        for c in confs:
            before = dbm.row(con, "SELECT id FROM conferences WHERE slug=?", (c["slug"],))
            upsert_conference(con, c, overwrite)
            if before:
                n_upd += 1
            else:
                n_new += 1
        if isinstance(data, dict):
            for s in data.get("sources", []):
                if not dbm.row(con, "SELECT id FROM sources WHERE url=?", (s["url"],)):
                    dbm.insert_row(con, "sources", {k: s.get(k) for k in ("name", "url", "notes")})
            for t in data.get("templates", []):
                if not dbm.row(con, "SELECT id FROM templates WHERE name=?", (t["name"],)):
                    dbm.insert_row(con, "templates", {k: t.get(k) for k in ("name", "subject", "body")})
        dbm.log_activity(con, "import", f"Loaded seed data: {n_new} new, {n_upd} updated")
    return {"new": n_new, "updated": n_upd, "total": len(confs)}
